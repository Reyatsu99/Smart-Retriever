from __future__ import annotations

import csv
import json
import shutil
import gc
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import docx
import openpyxl

from smart_retriever.benchmark import StubEmbedder
from smart_retriever.embeddings import EmbeddingBackend
from smart_retriever.indexer import build_index


DEFAULT_EVAL_WORK_DIR = Path(".benchmarks/v2/public/latest")


@dataclass(frozen=True)
class EvalQuery:
    query_id: str
    text: str


def run_public_evaluation(
    dataset_id: str,
    *,
    top_k: int = 10,
    max_queries: int | None = None,
    query_offset: int = 0,
    pool: str = "auto",
    pool_depth: int = 100,
    mode: str = "real",
    search_mode: str = "full",
    force: bool = False,
    work_dir: Path | str = DEFAULT_EVAL_WORK_DIR,
    keep_workspace: bool = True,
) -> dict[str, Any]:
    ir_datasets = _import_ir_datasets()

    print(f"Loading dataset: {dataset_id}")
    dataset, resolved_dataset_id = _load_eval_dataset(ir_datasets, dataset_id)
    work_dir = Path(work_dir)
    print(f"Working directory: {work_dir}")
    data_dir = work_dir / "data"
    index_dir = work_dir / "index"
    report_path = work_dir / "evaluation_report.json"
    run_path = work_dir / "run.json"
    mapping_path = work_dir / "doc_mapping.json"

    qrels_by_query = _load_qrels(dataset)
    selected_queries = _select_queries(dataset, qrels_by_query, max_queries=max_queries, query_offset=query_offset)
    if not selected_queries:
        raise ValueError(f"No judged queries available for dataset '{dataset_id}'.")

    resolved_pool = _resolve_pool(resolved_dataset_id, pool)
    doc_mapping = _materialize_corpus(
        dataset,
        data_dir=data_dir,
        queries=selected_queries,
        qrels_by_query=qrels_by_query,
        pool=resolved_pool,
        pool_depth=pool_depth,
    )

    embedder = _select_backends(mode)
    print("Building/Verifying index...")
    manifest = build_index(
        data_dir=data_dir,
        index_dir=index_dir,
        force=force,
        embedder=embedder,
    )
    print("Initializing SearchEngine...")
    from smart_retriever.search import SearchEngine

    engine = SearchEngine(index_dir=index_dir, embedder=embedder)
    print(f"Running {len(selected_queries)} queries (search_mode={search_mode})...")
    run_output = _run_queries(engine, selected_queries, doc_mapping, top_k=top_k, search_mode=search_mode)
    run = run_output["run"]
    avg_latency = run_output["avg_latency_ms"]
    
    print("Computing metrics...")
    metrics = _compute_metrics(run, qrels_by_query, top_k=top_k)
    metrics["avg_latency_ms"] = round(avg_latency, 2)
    
    print("Summarizing results...")
    summary = _summarize_queries(run, qrels_by_query, top_k=top_k)

    work_dir.mkdir(parents=True, exist_ok=True)
    mapping_path.write_text(json.dumps(doc_mapping, indent=2), encoding="utf-8")
    run_path.write_text(json.dumps(run, indent=2), encoding="utf-8")

    report = {
        "dataset_id": dataset_id,
        "resolved_dataset_id": resolved_dataset_id,
        "pool": resolved_pool,
        "pool_depth": pool_depth if resolved_pool == "scoreddocs" else None,
        "mode": mode,
        "search_mode": search_mode,
        "top_k": top_k,
        "query_offset": query_offset,
        "queries_evaluated": len(selected_queries),
        "documents_indexed": manifest.get("indexed_file_count", 0),
        "index_manifest": {
            "indexed_file_count": manifest.get("indexed_file_count", 0),
            "skipped_file_count": manifest.get("skipped_file_count", 0),
            "chunk_count": manifest.get("indexed_chunks", 0),
        },
        "metrics": metrics,
        "query_summary": summary,
        "paths": {
            "workspace": str(work_dir.resolve()),
            "run": str(run_path.resolve()),
            "mapping": str(mapping_path.resolve()),
            "report": str(report_path.resolve()),
        },
    }
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")

    if not keep_workspace:
        compact_report = dict(report)
        compact_report["paths"] = {"workspace": str(work_dir.resolve())}
        compact_report["workspace_deleted"] = True
        shutil.rmtree(work_dir)
        gc.collect()
        return compact_report
    gc.collect()
    return report


def _select_backends(mode: str) -> Any:
    if mode == "real":
        return EmbeddingBackend()
    if mode == "stub":
        return StubEmbedder()
    raise ValueError("mode must be 'real' or 'stub'")


def _import_ir_datasets() -> Any:
    try:
        import ir_datasets
    except ImportError as exc:
        raise ImportError(
            "ir_datasets is required for public benchmark evaluation. Install it with `pip install ir_datasets`."
        ) from exc
    return ir_datasets


def _import_ir_measures() -> Any:
    try:
        import ir_measures
    except ImportError as exc:
        raise ImportError(
            "ir_measures is required for standard evaluation metrics. Install it with `pip install ir_measures`."
        ) from exc
    return ir_measures


def _load_eval_dataset(ir_datasets: Any, dataset_id: str) -> tuple[Any, str]:
    dataset = ir_datasets.load(dataset_id)
    if hasattr(dataset, "qrels_iter"):
        return dataset, dataset_id
    for suffix in ("/test", "/dev/search", "/dev", "/train"):
        candidate_id = f"{dataset_id}{suffix}"
        try:
            candidate = ir_datasets.load(candidate_id)
        except KeyError:
            continue
        if hasattr(candidate, "qrels_iter"):
            return candidate, candidate_id
    raise ValueError(
        f"Dataset '{dataset_id}' does not provide qrels directly. "
        "Use a judged split such as '/test', '/dev', or '/dev/search'."
    )


def _resolve_pool(dataset_id: str, pool: str) -> str:
    if pool != "auto":
        return pool
    if dataset_id.startswith("msmarco-passage/"):
        return "scoreddocs"
    return "full"


def _load_qrels(dataset: Any) -> dict[str, dict[str, int]]:
    qrels_by_query: dict[str, dict[str, int]] = defaultdict(dict)
    for qrel in dataset.qrels_iter():
        relevance = int(getattr(qrel, "relevance", 0))
        if relevance <= 0:
            continue
        qrels_by_query[qrel.query_id][qrel.doc_id] = relevance
    return dict(qrels_by_query)


def _select_queries(
    dataset: Any,
    qrels_by_query: dict[str, dict[str, int]],
    *,
    max_queries: int | None,
    query_offset: int,
) -> list[EvalQuery]:
    selected: list[EvalQuery] = []
    skipped = 0
    for query in dataset.queries_iter():
        if query.query_id not in qrels_by_query:
            continue
        if skipped < query_offset:
            skipped += 1
            continue
        selected.append(EvalQuery(query_id=query.query_id, text=query.text))
        if max_queries is not None and len(selected) >= max_queries:
            break
    return selected


def _materialize_corpus(
    dataset: Any,
    *,
    data_dir: Path,
    queries: list[EvalQuery],
    qrels_by_query: dict[str, dict[str, int]],
    pool: str,
    pool_depth: int,
) -> dict[str, dict[str, str]]:
    doc_mapping: dict[str, dict[str, str]] = {}
    data_dir.mkdir(parents=True, exist_ok=True)
    if pool == "full":
        for ordinal, doc in enumerate(dataset.docs_iter()):
            relative_path = _relative_doc_path(ordinal)
            _write_doc_file(data_dir / relative_path, _doc_text(doc))
            doc_mapping[str(getattr(doc, "doc_id"))] = {
                "relative_path": relative_path,
                "file_name": Path(relative_path).name,
            }
        return doc_mapping

    doc_ids = _candidate_doc_ids(dataset, queries, qrels_by_query, pool=pool, pool_depth=pool_depth)
    
    # Use docs_iter to avoid loading the entire docs_store into memory
    count = 0
    for doc in dataset.docs_iter():
        doc_id = str(getattr(doc, "doc_id"))
        if doc_id in doc_ids:
            relative_path = _relative_doc_path(count)
            _write_doc_file(data_dir / relative_path, _doc_text(doc))
            doc_mapping[doc_id] = {
                "relative_path": relative_path,
                "file_name": Path(relative_path).name,
            }
            count += 1
            if len(doc_mapping) == len(doc_ids):
                break
    return doc_mapping


def _candidate_doc_ids(
    dataset: Any,
    queries: list[EvalQuery],
    qrels_by_query: dict[str, dict[str, int]],
    *,
    pool: str,
    pool_depth: int,
) -> set[str]:
    query_ids = {query.query_id for query in queries}
    doc_ids: set[str] = set()
    if pool == "qrels":
        for query_id in query_ids:
            doc_ids.update(qrels_by_query.get(query_id, {}))
        return doc_ids
    if pool == "scoreddocs":
        if not hasattr(dataset, "scoreddocs_iter"):
            raise ValueError("Selected dataset does not expose scoreddocs; use --pool full or --pool qrels instead.")
        counts: dict[str, int] = defaultdict(int)
        for scored_doc in dataset.scoreddocs_iter():
            query_id = scored_doc.query_id
            if query_id not in query_ids or counts[query_id] >= pool_depth:
                continue
            doc_ids.add(scored_doc.doc_id)
            counts[query_id] += 1
        for query_id in query_ids:
            doc_ids.update(qrels_by_query.get(query_id, {}))
        return doc_ids
    raise ValueError("pool must be one of 'auto', 'full', 'scoreddocs', or 'qrels'")


def _relative_doc_path(ordinal: int) -> str:
    ext = ".txt"
    bucket = f"docs_{ordinal // 1000:05d}"
    return str(Path(bucket) / f"doc_{ordinal:07d}{ext}")


def _write_doc_file(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    ext = path.suffix.lower()
    if ext == ".docx":
        document = docx.Document()
        document.add_paragraph(text)
        document.save(path)
    elif ext == ".xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        for idx, line in enumerate(text.split("\n"), start=1):
            ws.cell(row=idx, column=1, value=line)
        wb.save(path)
    elif ext == ".csv":
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            for line in text.split("\n"):
                writer.writerow([line])
    else:
        path.write_text(text, encoding="utf-8")


def _doc_text(doc: Any) -> str:
    parts: list[str] = []
    field_names = getattr(doc, "_fields", ())
    if field_names:
        for field_name in field_names:
            if field_name in {"doc_id", "id", "url"}:
                continue
            value = getattr(doc, field_name, None)
            if value is None:
                continue
            text = str(value).strip()
            if text:
                parts.append(text)
    elif hasattr(doc, "text"):
        text = str(getattr(doc, "text")).strip()
        if text:
            parts.append(text)
    return "\n".join(parts)


def _run_queries(engine: Any, queries: list[EvalQuery], doc_mapping: dict[str, dict[str, str]], *, top_k: int, search_mode: str = "full") -> dict[str, Any]:
    import time
    doc_id_by_path = {payload["relative_path"]: doc_id for doc_id, payload in doc_mapping.items()}
    run_results: dict[str, list[dict[str, Any]]] = {}
    latencies = []
    
    for query in queries:
        start_time = time.perf_counter()
        results = engine.search(query.text, top_k=top_k, search_mode=search_mode)
        end_time = time.perf_counter()
        latencies.append((end_time - start_time) * 1000) # ms
        
        converted: list[dict[str, Any]] = []
        for rank, result in enumerate(results, start=1):
            doc_id = doc_id_by_path.get(result["relative_path"])
            if doc_id is None:
                continue
            converted.append(
                {
                    "doc_id": doc_id,
                    "rank": rank,
                    "score": result["score"],
                    "relative_path": result["relative_path"],
                }
            )
        run_results[query.query_id] = converted
        
    return {
        "run": run_results,
        "avg_latency_ms": sum(latencies) / len(latencies) if latencies else 0.0
    }


def _compute_metrics(run: dict[str, list[dict[str, Any]]], qrels_by_query: dict[str, dict[str, int]], *, top_k: int) -> dict[str, float]:
    ir_measures = _import_ir_measures()
    measures = [
        ir_measures.parse_measure(f"nDCG@{top_k}"),
        ir_measures.parse_measure(f"RR@{top_k}"),
        ir_measures.parse_measure(f"R@{top_k}"),
        ir_measures.parse_measure(f"P@{top_k}"),
        ir_measures.parse_measure("AP"),
    ]
    qrels = {query_id: {doc_id: relevance for doc_id, relevance in docrels.items()} for query_id, docrels in qrels_by_query.items()}
    scored_run = {query_id: {item["doc_id"]: float(item["score"]) for item in ranked_docs} for query_id, ranked_docs in run.items()}
    aggregate = ir_measures.calc_aggregate(measures, qrels, scored_run)
    return {
        f"ndcg@{top_k}": round(float(aggregate[measures[0]]), 4),
        f"mrr@{top_k}": round(float(aggregate[measures[1]]), 4),
        f"recall@{top_k}": round(float(aggregate[measures[2]]), 4),
        f"precision@{top_k}": round(float(aggregate[measures[3]]), 4),
        "map": round(float(aggregate[measures[4]]), 4),
        f"success@{top_k}": round(_success_at_k(run, qrels_by_query, top_k=top_k), 4),
    }


def _summarize_queries(run: dict[str, list[dict[str, Any]]], qrels_by_query: dict[str, dict[str, int]], *, top_k: int) -> dict[str, Any]:
    misses: list[str] = []
    first_hit_ranks: list[int] = []
    for query_id, ranked_docs in run.items():
        relevant_ids = {doc_id for doc_id, rel in qrels_by_query.get(query_id, {}).items() if rel > 0}
        hit_rank = None
        for rank, item in enumerate(ranked_docs[:top_k], start=1):
            if item["doc_id"] in relevant_ids:
                hit_rank = rank
                break
        if hit_rank is None:
            misses.append(query_id)
        else:
            first_hit_ranks.append(hit_rank)
    return {
        "queries_with_hit": len(first_hit_ranks),
        "queries_without_hit": len(misses),
        "avg_first_hit_rank": round(sum(first_hit_ranks) / len(first_hit_ranks), 3) if first_hit_ranks else None,
        "sample_missed_query_ids": misses[:10],
    }


def _success_at_k(run: dict[str, list[dict[str, Any]]], qrels_by_query: dict[str, dict[str, int]], *, top_k: int) -> float:
    hits = 0
    query_count = 0
    for query_id, ranked_docs in run.items():
        relevant_ids = {doc_id for doc_id, rel in qrels_by_query.get(query_id, {}).items() if rel > 0}
        if not relevant_ids:
            continue
        query_count += 1
        ranked_doc_ids = [item["doc_id"] for item in ranked_docs[:top_k]]
        hits += int(any(doc_id in relevant_ids for doc_id in ranked_doc_ids))
    return hits / query_count if query_count else 0.0
